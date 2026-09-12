import os
import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font
from docx import Document
from number_to_words import doc_so_thanh_chu

def format_money(val):
    """Chuyển đổi số tiền thành chuỗi có dấu chấm phân cách hàng nghìn (VD: 2651658 -> 2.651.658)"""
    try:
        num = int(str(val).replace('.', '').replace(',', ''))
        return f"{num:,}".replace(',', '.')
    except:
        return val

def generate_outputs(data, template_dir="templates", output_dir="output"):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    full_output_dir = os.path.join(base_dir, output_dir)
    os.makedirs(full_output_dir, exist_ok=True)
    
    # Xử lý thông tin kỳ cước từ dữ liệu (VD: T08-2026)
    thang_str = "08"
    nam_str = "2026"
    try:
        parts = data.get('tu_ngay', '').split('/')
        if len(parts) == 3:
            thang_str = parts[1]
            nam_str = parts[2]
    except:
        pass
    ky_cuoc = f"T{thang_str}-{nam_str}"
    
    # Làm sạch tên công ty để tạo tên file an toàn (không chứa ký tự đặc biệt)
    ten_kh_raw = data.get('ten_khach_hang', 'KhachHang')
    ten_kh_safe = "".join([c if c.isalnum() or c in (' ', '_', '-') else '' for c in ten_kh_raw]).strip().replace(' ', '_')
    
    # Tạo hậu tố chung cho tên file: _Ten_Cong_Ty-Ky_Cuoc
    file_suffix = f"{ten_kh_safe}-{ky_cuoc}"

    # ==========================================
    # 1. XỬ LÝ XUẤT BẢNG KÊ EXCEL
    # ==========================================
    excel_template = os.path.join(base_dir, template_dir, "bang_ke.xlsx")
    excel_output = os.path.join(full_output_dir, f"BangKe_{file_suffix}.xlsx")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    if os.path.exists(excel_template):
        wb = openpyxl.load_workbook(excel_template)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bảng kê"
        ws['A4'] = "STT"
        ws['B4'] = "Số TB"
        ws['C4'] = "Tên KH"
        ws['D4'] = "Địa chỉ TT"
        ws['E4'] = "Số ĐT LH"
        ws['F4'] = "Mã Số Thuế"
        ws['G4'] = "Tổng cộng"
        for col in range(1, 8):
            ws.cell(row=4, column=col).font = Font(name='Times New Roman', size=11, bold=True)
            ws.cell(row=4, column=col).alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(row=4, column=col).border = thin_border

    ws['A1'] = data['ten_khach_hang']
    ws['A2'] = f"Danh Sách Cước từ {data['tu_ngay']} - {data['den_ngay']}"
    
    start_row = 5
    tong_cong_val = 0
    
    for idx, item in enumerate(data['chi_tiet']):
        row = start_row + idx
        ws[f'A{row}'] = item['stt']
        ws[f'B{row}'] = str(item['so_thue_bao'])
        ws[f'C{row}'] = data['ten_khach_hang']
        ws[f'D{row}'] = data['dia_chi']
        ws[f'E{row}'] = "773998999"
        ws[f'F{row}'] = data['mst']
        
        cong_str = format_money(item['cong'])
        ws[f'G{row}'] = cong_str
        
        tong_cong_val += int(str(item['cong']).replace('.', '').replace(',', ''))
        
        for col in range(1, 8):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.font = Font(name='Times New Roman', size=11)
            if col in [1, 2, 5, 6, 7]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            
    row_tong_cong = start_row + len(data['chi_tiet'])
    ws[f'A{row_tong_cong}'] = "Tổng cộng"
    ws[f'G{row_tong_cong}'] = format_money(tong_cong_val)
    
    for col in range(1, 8):
        cell = ws.cell(row=row_tong_cong, column=col)
        cell.border = thin_border
        cell.font = Font(name='Times New Roman', size=11, bold=True)
    
    ws.merge_cells(start_row=row_tong_cong, start_column=1, end_row=row_tong_cong, end_column=6)
    ws.cell(row=row_tong_cong, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=row_tong_cong, column=7).alignment = Alignment(horizontal='center', vertical='center')
    
    wb.save(excel_output)

    # ==========================================
    # 2. XỬ LÝ XUẤT GIẤY ĐỀ NGHỊ THANH TOÁN WORD
    # ==========================================
    docx_template = os.path.join(base_dir, template_dir, "DNTT.docx")
    docx_output = os.path.join(full_output_dir, f"DNTT_{file_suffix}.docx")
    
    so_tien_chu = doc_so_thanh_chu(data['tong_cong'])
    formatted_tong_cong = format_money(data['tong_cong'])
    
    ky_cuoc_slash = f"T{thang_str}/{nam_str}"
    ngay_lap_hd = data.get('ngay_hoa_don', f"05/{thang_str}/{nam_str}")

    if os.path.exists(docx_template):
        doc = Document(docx_template)
        
        def process_paragraph(p):
            full_text = p.text
            if "V/v Đề nghị thanh toán" in full_text or "MobiFone T" in full_text:
                p.text = f"V/v Đề nghị thanh toán cước phí viễn thông\nMobiFone {ky_cuoc_slash}."
            elif "An Giang, ngày" in full_text:
                p.text = f"An Giang, ngày {ngay_lap_hd}"
            elif "Kính gửi:" in full_text:
                p.text = f"Kính gửi: {data['ten_khach_hang']}"
            elif "Chúng tôi xin thông báo đến Quý Công ty" in full_text or "cước điện thoại di động" in full_text:
                p.text = (
                    f"Chúng tôi xin thông báo đến Quý Công ty thông tin cước điện thoại di động của Quý Công ty, "
                    f"tháng {thang_str}/{nam_str} là: {formatted_tong_cong} (Bằng chữ : {so_tien_chu})"
                )
            elif "MobiFone An Giang_" in full_text:
                p.text = f"MobiFone An Giang_{data['ten_khach_hang']}_{data['mst']}_Cước viễn thông tháng {ky_cuoc_slash}"

        # 1. Quét phần Header của tất cả các section
        for section in doc.sections:
            header = section.header
            for p in header.paragraphs:
                process_paragraph(p)
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            process_paragraph(p)

        # 2. Quét phần thân văn bản (Body)
        for p in doc.paragraphs:
            process_paragraph(p)
            
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        process_paragraph(p)
                        
        doc.save(docx_output)
    else:
        # Fallback
        doc = Document()
        doc.add_heading("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", level=2)
        doc.add_paragraph("Độc lập - Tự do - Hạnh phúc\n")
        doc.add_heading("GIẤY ĐỀ NGHỊ THANH TOÁN", level=1)
        doc.add_paragraph(f"Kính gửi: Ban Giám đốc / Phòng Kế toán")
        doc.add_paragraph(f"- Đơn vị: {data['ten_khach_hang']}")
        doc.add_paragraph(f"- Nội dung thanh toán: Thanh toán cước viễn thông theo hóa đơn số {data['so_hoa_don']} ngày {data['ngay_hoa_don']}")
        doc.add_paragraph(f"- Số tiền đề nghị thanh toán: {formatted_tong_cong} VNĐ")
        doc.add_paragraph(f"- Bằng chữ: {so_tien_chu}")
        doc.save(docx_output)

    return excel_output, docx_output